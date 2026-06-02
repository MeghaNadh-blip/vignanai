import os
import random
import smtplib
import shutil
import google.generativeai as genai
from datetime import datetime, timedelta
from email.message import EmailMessage
from email_validator import validate_email, EmailNotValidError
from uuid import uuid4
from functools import wraps
from io import BytesIO
from textwrap import wrap
import certifi
import bcrypt
from bson import ObjectId
from dotenv import load_dotenv
from flask import Flask, Response, flash, redirect, render_template, request, session, url_for
from markupsafe import Markup, escape
from groq import Groq
import google.generativeai as genai
from pymongo import MongoClient
from pymongo.errors import PyMongoError
from flask_wtf.csrf import CSRFProtect
from email_validator import validate_email, EmailNotValidError


load_dotenv()

app = Flask(__name__)

# Security configuration
app.secret_key = os.getenv("SECRET_KEY")
if not app.secret_key:
    app.secret_key = "dev-secret-change-me"
    print("WARNING: SECRET_KEY is not set. Add a strong SECRET_KEY in your .env before deployment.")

app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024  # 16 MB upload limit
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.getenv("FLASK_ENV", "development").lower() == "production"


COMMON_EMAIL_TYPOS = {
    "gamil.com": "gmail.com",
    "gmial.com": "gmail.com",
    "gmai.com": "gmail.com",
    "gmail.co": "gmail.com",
    "yaho.com": "yahoo.com",
    "yahoo.co": "yahoo.com",
    "outlok.com": "outlook.com",
    "hotmial.com": "hotmail.com",
}


def normalize_and_validate_email(raw_email):
    """Return a normalized email or raise ValueError with a student-friendly message."""
    email = (raw_email or "").strip().lower()
    if not email:
        raise ValueError("Please enter an email address.")

    domain = email.split("@")[-1] if "@" in email else ""
    if domain in COMMON_EMAIL_TYPOS:
        raise ValueError(f"Email domain looks wrong. Did you mean @{COMMON_EMAIL_TYPOS[domain]}?")

    try:
        # check_deliverability=False avoids network delays, while still blocking invalid formats like m@gamil or abc@.
        valid = validate_email(email, check_deliverability=False)
        return valid.normalized.lower()
    except EmailNotValidError as exc:
        raise ValueError(f"Please enter a valid email address. {exc}") from exc

csrf = CSRFProtect(app)

mongo_uri = os.getenv("MONGO_URI", "mongodb://localhost:27017/").strip()

# MongoDB connection
# Local development uses mongodb://localhost.
# Render/Atlas must use mongodb+srv:// and needs TLS options to avoid SSL handshake issues.
# connect=False prevents Gunicorn from opening MongoDB before worker fork.
mongo_options = {
    "connect": False,
    "serverSelectionTimeoutMS": int(os.getenv("MONGO_TIMEOUT_MS", "10000")),
    "connectTimeoutMS": int(os.getenv("MONGO_TIMEOUT_MS", "10000")),
    "socketTimeoutMS": int(os.getenv("MONGO_TIMEOUT_MS", "10000")),
}

if mongo_uri.startswith("mongodb+srv://"):
    mongo_options.update({
        "tls": True,
        "tlsCAFile": certifi.where(),
        # Render + Atlas can sometimes fail certificate handshake; this keeps login working.
        # For stricter production security, set MONGO_TLS_ALLOW_INVALID=false after confirming Atlas TLS works.
        "tlsAllowInvalidCertificates": os.getenv("MONGO_TLS_ALLOW_INVALID", "true").lower() == "true",
    })

mongo_client = MongoClient(mongo_uri, **mongo_options)
db = mongo_client[os.getenv("MONGO_DB_NAME", "vignan_1")]
users_collection = db["users"]
feedback_collection = db["feedback"]
question_generator_collection = db["question_generator"]
generated_outputs_collection = db["generated_outputs"]


def save_generated_output(output_type, content, metadata=None):
    """Store generated AI output in MongoDB and keep only its ID in the Flask session."""
    record = {
        "user_email": session.get("email"),
        "username": session.get("name", "Student"),
        "output_type": output_type,
        "content": content,
        "metadata": metadata or {},
        "created_at": datetime.utcnow(),
    }
    result = generated_outputs_collection.insert_one(record)
    return str(result.inserted_id)


def get_generated_output(output_id, expected_type=None):
    """Read one generated output for the logged-in user from MongoDB."""
    if not output_id:
        return None
    try:
        query = {"_id": ObjectId(output_id), "user_email": session.get("email")}
    except Exception:
        return None
    if expected_type:
        query["output_type"] = expected_type
    return generated_outputs_collection.find_one(query)


def get_generated_content_from_session(session_key, expected_type=None):
    record = get_generated_output(session.get(session_key), expected_type)
    return record.get("content") if record else None


@app.after_request
def add_security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    return response


@app.template_filter("ai_format")
def ai_format(value):
    """Render AI text as safe, readable HTML cards without external markdown packages."""
    if value is None:
        return ""

    import re

    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return ""

    def inline_format(raw):
        safe = str(escape(raw))
        safe = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", safe)
        safe = re.sub(r"`([^`]+)`", r"<code class=\"inline-code\">\1</code>", safe)
        return safe

    html = []
    list_open = False
    code_open = False
    code_lines = []

    def close_list():
        nonlocal list_open
        if list_open:
            html.append("</ul>")
            list_open = False

    def close_code():
        nonlocal code_open, code_lines
        if code_open:
            escaped_code = escape("\n".join(code_lines))
            html.append(f'<pre class="ai-code-block"><code>{escaped_code}</code></pre>')
            code_lines = []
            code_open = False

    for raw_line in text.split("\n"):
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if code_open:
                close_code()
            else:
                close_list()
                code_open = True
                code_lines = []
            continue

        if code_open:
            code_lines.append(line)
            continue

        if not stripped:
            close_list()
            continue

        lower = stripped.lower()
        heading_markers = (
            "based on uploaded material",
            "additional ai knowledge",
            "ai knowledge",
            "uploaded material",
            "key points",
            "explanation",
            "answer",
            "summary",
            "examples",
            "important questions",
            "short notes",
            "mcqs",
            "viva questions",
            "source",
        )

        if stripped.startswith("#"):
            close_list()
            title = stripped.lstrip("#").strip()
            html.append(f'<h3 class="ai-section-title">{inline_format(title)}</h3>')
            continue

        if any(lower.startswith(marker) for marker in heading_markers) and len(stripped) <= 80:
            close_list()
            title = stripped.rstrip(":")
            badge_class = " ai-source-title" if "source" in lower or "uploaded" in lower or "ai knowledge" in lower else ""
            html.append(f'<h3 class="ai-section-title{badge_class}">{inline_format(title)}</h3>')
            continue

        if stripped.startswith(("- ", "• ", "* ")):
            if not list_open:
                html.append('<ul class="ai-list">')
                list_open = True
            html.append(f'<li>{inline_format(stripped[2:].strip())}</li>')
            continue

        numbered = re.match(r"^(\d+)[\.)]\s+(.*)", stripped)
        if numbered:
            close_list()
            html.append(
                '<div class="ai-numbered-item">'
                f'<span class="ai-number-pill">{numbered.group(1)}</span>'
                f'<p>{inline_format(numbered.group(2))}</p>'
                '</div>'
            )
            continue

        close_list()
        if stripped.startswith(("📄", "🤖", "✅", "⚠️", "❌", "🧠", "📌")):
            html.append(f'<p class="ai-highlight-line">{inline_format(stripped)}</p>')
        else:
            html.append(f'<p>{inline_format(stripped)}</p>')

    close_code()
    close_list()
    return Markup("\n".join(html))

# AI provider configuration
# Default is Groq because this is the active provider in the current app.
# Gemini can be enabled later by setting AI_PROVIDER=gemini and adding GOOGLE_API_KEY/GEMINI_API_KEY.
AI_PROVIDER = os.getenv("AI_PROVIDER", "groq").strip().lower()

groq_api_key = os.getenv("GROQ_API_KEY")
groq_client = Groq(api_key=groq_api_key) if groq_api_key else None
GROQ_MODEL = os.getenv("GROQ_MODEL", "llama-3.1-8b-instant")

gemini_api_key = os.getenv("GOOGLE_API_KEY") or os.getenv("GEMINI_API_KEY")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
if gemini_api_key:
    genai.configure(api_key=gemini_api_key)


VFSTR_SYSTEM_PROMPT = """
You are VFSTR Academic AI Assistant for Vignan's Foundation for Science, Technology & Research (VFSTR).
Help students with academics, coding, assignments, notes, quizzes, planning, and project guidance.
First use uploaded/user-provided material when available. If it is incomplete, clearly add helpful AI general knowledge.
Mention whether the answer is based on uploaded material, AI knowledge, or both when relevant.
Be clear, concise, student-friendly, and practical.
"""


def friendly_ai_error(error):
    """Return a user-friendly error instead of exposing raw API traceback text."""
    message = str(error)
    lower = message.lower()
    if "groq_api_key" in lower or "api key" in lower or "invalid api" in lower or "401" in lower:
        return (
            "AI service is not ready. Please check your AI API key in the .env file "
            "(GROQ_API_KEY for Groq or GOOGLE_API_KEY/GEMINI_API_KEY for Gemini), save it, and restart Flask."
        )
    if "rate" in lower or "429" in lower:
        return "AI service is busy or rate-limited. Please try again after a short break."
    if "connection" in lower or "timeout" in lower or "network" in lower:
        return "Could not connect to the AI service. Please check internet connection and try again."
    return "Something went wrong while generating the AI response. Please try again."




def send_reset_email(to_email, otp_code):
    """Send a password reset OTP by email using SMTP settings from .env."""
    mail_username = os.getenv("MAIL_USERNAME")
    mail_password = os.getenv("MAIL_PASSWORD")
    mail_server = os.getenv("MAIL_SERVER", "smtp.gmail.com")
    mail_port = int(os.getenv("MAIL_PORT", "587"))

    if not mail_username or not mail_password:
        return False, "Email service is not configured. Add MAIL_USERNAME and MAIL_PASSWORD in .env."

    msg = EmailMessage()
    msg["Subject"] = "Vignan AI password reset code"
    msg["From"] = mail_username
    msg["To"] = to_email
    msg.set_content(
        f"""Hello,

Your Vignan AI password reset code is: {otp_code}

This code is valid for 10 minutes. If you did not request this, you can ignore this email.

Vignan AI
"""
    )

    try:
        with smtplib.SMTP(mail_server, mail_port, timeout=20) as server:
            server.starttls()
            server.login(mail_username, mail_password)
            server.send_message(msg)
        return True, "Reset code sent successfully."
    except Exception as exc:
        return False, f"Could not send email: {exc}"

def is_allowed_upload(filename):
    allowed = (".pdf", ".docx", ".pptx", ".txt", ".png", ".jpg", ".jpeg", ".webp", ".bmp")
    return bool(filename) and filename.lower().endswith(allowed)


def extract_text_from_pdf(uploaded_file):
    """Extract text from normal PDFs and OCR-based/scanned PDFs.

    First tries PyPDF2 for text-based PDFs. If that gives no text, it falls
    back to rendering pages with pypdfium2 and reading them with Tesseract OCR.
    """
    import PyPDF2

    uploaded_file.seek(0)
    pdf_bytes = uploaded_file.read()

    text_parts = []
    try:
        pdf_reader = PyPDF2.PdfReader(BytesIO(pdf_bytes))
        for page in pdf_reader.pages:
            page_text = page.extract_text()
            if page_text and page_text.strip():
                text_parts.append(page_text.strip())
    except Exception:
        pass

    extracted_text = "\n".join(text_parts).strip()
    if extracted_text:
        return extracted_text

    # OCR fallback for scanned/image-only PDFs.
    try:
        import pypdfium2 as pdfium
        import pytesseract
        from pytesseract import TesseractNotFoundError
    except ImportError as exc:
        raise ValueError(
            "This PDF looks scanned/image-based. Install pypdfium2 and pytesseract for OCR support."
        ) from exc

    if shutil.which("tesseract") is None:
        raise ValueError(
            "This PDF looks scanned/image-based, but Tesseract OCR is not installed on this system. "
            "Install Tesseract or upload a text-based PDF/DOCX/TXT file."
        )

    ocr_parts = []
    try:
        pdf = pdfium.PdfDocument(pdf_bytes)
        for page_index in range(len(pdf)):
            page = pdf[page_index]
            bitmap = page.render(scale=2).to_pil()
            page_text = pytesseract.image_to_string(bitmap)
            if page_text and page_text.strip():
                ocr_parts.append(page_text.strip())
    except TesseractNotFoundError as exc:
        raise ValueError(
            "Tesseract OCR is not installed or not available in PATH. Upload a text-based file or install Tesseract."
        ) from exc
    except Exception as exc:
        raise ValueError(
            "Could not read this PDF. If it is scanned, please install Tesseract OCR on your system."
        ) from exc

    return "\n".join(ocr_parts).strip()


def extract_text_from_docx(uploaded_file):
    from docx import Document
    uploaded_file.seek(0)
    document = Document(uploaded_file)
    return "\n".join(p.text for p in document.paragraphs if p.text.strip()).strip()


def extract_text_from_pptx(uploaded_file):
    from pptx import Presentation
    uploaded_file.seek(0)
    presentation = Presentation(uploaded_file)
    text_parts = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        if slide_text:
            text_parts.append(f"Slide {slide_number}:\\n" + "\n".join(slide_text))
    return "\n\n".join(text_parts).strip()


def extract_text_from_image(uploaded_file):
    from PIL import Image
    try:
        import pytesseract
        from pytesseract import TesseractNotFoundError
    except ImportError as exc:
        raise ValueError("Image OCR needs pytesseract. Install pytesseract or upload text-based material.") from exc

    if shutil.which("tesseract") is None:
        raise ValueError(
            "Image OCR needs Tesseract installed on your system. Upload PDF/DOCX/TXT material or install Tesseract."
        )

    uploaded_file.seek(0)
    image = Image.open(uploaded_file)
    try:
        return pytesseract.image_to_string(image).strip()
    except TesseractNotFoundError as exc:
        raise ValueError("Tesseract OCR is not installed or not available in PATH.") from exc


def extract_text_from_file(uploaded_file):
    if not uploaded_file or not uploaded_file.filename:
        raise ValueError("Please upload a file.")
    if not is_allowed_upload(uploaded_file.filename):
        raise ValueError("Supported files: PDF, DOCX, PPTX, TXT, PNG, JPG, JPEG, WEBP, BMP.")
    filename = uploaded_file.filename.lower()
    if filename.endswith(".txt"):
        uploaded_file.seek(0)
        return uploaded_file.read().decode("utf-8", errors="ignore").strip()
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    if filename.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    if filename.endswith(".pptx"):
        return extract_text_from_pptx(uploaded_file)
    if filename.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp")):
        return extract_text_from_image(uploaded_file)
    raise ValueError("Supported files: PDF, DOCX, PPTX, TXT, PNG, JPG, JPEG, WEBP, BMP.")


def create_docx_response(title, body, filename):
    from docx import Document
    buffer = BytesIO()
    document = Document()
    document.add_heading(title, 0)
    for block in body.split("\n"):
        line = block.strip()
        if not line:
            document.add_paragraph("")
        elif line.startswith(("- ", "• ")):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    document.save(buffer)
    buffer.seek(0)
    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}.docx"},
    )


def create_pdf_response(title, body, filename):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter
    pdf.setTitle(title)
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(72, height - 72, title)
    y = height - 105
    pdf.setFont("Helvetica", 10)
    for raw_line in body.split("\n"):
        for line in wrap(raw_line, width=95) or [""]:
            if y < 72:
                pdf.showPage()
                pdf.setFont("Helvetica", 10)
                y = height - 72
            pdf.drawString(72, y, line)
            y -= 15
        if raw_line.strip() == "":
            y -= 5
    pdf.save()
    buffer.seek(0)
    return Response(
        buffer.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}.pdf"},
    )


def ask_ai(system_message, prompt):
    """Generate AI text using the configured provider: Groq by default, Gemini optionally."""
    if AI_PROVIDER == "gemini":
        if not gemini_api_key:
            raise ValueError("GOOGLE_API_KEY or GEMINI_API_KEY is missing in the .env file.")
        model = genai.GenerativeModel(GEMINI_MODEL)
        response = model.generate_content(f"{system_message}\n\nUser request:\n{prompt}")
        return (response.text or "").strip()

    if groq_client is None:
        raise ValueError("GROQ_API_KEY is missing in the .env file.")
    response = groq_client.chat.completions.create(
        model=GROQ_MODEL,
        messages=[{"role": "system", "content": system_message}, {"role": "user", "content": prompt}],
    )
    return response.choices[0].message.content


def ask_groq(system_message, prompt):
    """Backward-compatible wrapper. Existing routes can call this safely."""
    return ask_ai(system_message, prompt)




def generate_chat_answer(active_chat, question, uploaded_context="", uploaded_name=""):
    """Generate a chat answer using conversation plus optional uploaded file/image context.

    Important behavior:
    - If a file/image is uploaded, the answer must use that material first.
    - Only add general AI knowledge when the uploaded material is incomplete.
    - The response must clearly separate uploaded-material answers from extra AI knowledge.
    """
    chat_system_prompt = VFSTR_SYSTEM_PROMPT + """

Chat file rule:
When the user uploads a file or image, answer primarily from the uploaded content.
Use these headings exactly when a file is provided:
📄 Based on Uploaded Material
🤖 Additional AI Knowledge
If the uploaded material fully answers the question, keep Additional AI Knowledge very short or say it is not needed.
Do not pretend unsupported information came from the upload.
If the file text is unclear or OCR is weak, say that clearly and then help with careful AI knowledge.
"""

    groq_messages = [{"role": "system", "content": chat_system_prompt}]

    for msg in active_chat.get("messages", [])[-12:]:
        role = msg.get("role", "user")
        text_value = msg.get("text", "")
        if msg.get("attachment_name"):
            text_value += f"\n[Previous uploaded file: {msg.get('attachment_name')}]"
        if role in ["user", "assistant"] and text_value:
            groq_messages.append({"role": role, "content": text_value})

    clean_question = question.strip() or "Please read the uploaded file/image and explain the important content clearly."

    if uploaded_context:
        final_question = f"""
User question:
{clean_question}

Uploaded file name:
{uploaded_name or "uploaded material"}

Uploaded file/image extracted text:
{uploaded_context[:15000]}

Answer instructions:
1. Start with the heading: 📄 Based on Uploaded Material
2. Answer using the uploaded content first.
3. If useful, add the heading: 🤖 Additional AI Knowledge
4. Keep the answer student-friendly, clear, and practical.
""".strip()
    else:
        final_question = clean_question

    groq_messages.append({"role": "user", "content": final_question})

    if AI_PROVIDER == "gemini":
        # Gemini does not use the Groq/OpenAI chat format directly, so flatten recent context safely.
        flattened_context = "\n\n".join(f"{m['role'].title()}: {m['content']}" for m in groq_messages[1:])
        return ask_ai(chat_system_prompt, flattened_context)

    if groq_client is None:
        raise ValueError("GROQ_API_KEY is missing in the .env file.")
    response = groq_client.chat.completions.create(model=GROQ_MODEL, messages=groq_messages)
    return response.choices[0].message.content


def find_user_chat(user_email, chat_id):
    """Return user, conversations, selected chat, and selected index."""
    user = users_collection.find_one({"email": user_email})
    chat_conversations = user.get("chat_conversations", []) if user else []
    for index, conversation in enumerate(chat_conversations):
        if conversation.get("id") == chat_id:
            return user, chat_conversations, conversation, index
    return user, chat_conversations, None, -1

def login_required(view_function):
    @wraps(view_function)
    def wrapper(*args, **kwargs):
        if "user_id" not in session:
            flash("Please login first.", "error")
            return redirect(url_for("login"))
        return view_function(*args, **kwargs)

    return wrapper


def role_required(role):
    def decorator(view_function):
        @wraps(view_function)
        def wrapper(*args, **kwargs):
            if "user_id" not in session:
                flash("Please login first.", "error")
                return redirect(url_for("login"))

            if session.get("role") != role:
                flash("You do not have permission to open that page.", "error")
                if session.get("role") == "admin":
                    return redirect(url_for("admin"))
                return redirect(url_for("dashboard"))

            return view_function(*args, **kwargs)

        return wrapper

    return decorator



@app.route("/health/db")
def health_db():
    try:
        mongo_client.admin.command("ping")
        return {"status": "ok", "database": db.name}, 200
    except Exception as error:
        app.logger.error("MongoDB health check failed: %s", error)
        return {"status": "error", "message": str(error)}, 500

@app.route("/")
def home():
    if session.get("role") == "admin":
        return redirect(url_for("admin"))
    if session.get("role") == "student":
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        raw_email = request.form.get("email", "")
        try:
            email = normalize_and_validate_email(raw_email)
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("register"))
        password = request.form.get("password", "")
        role = request.form.get("role", "student").strip().lower()
        admin_code = request.form.get("admin_code", "").strip()

        if role not in ["student", "admin"]:
            role = "student"

        if not name or not email or not password:
            flash("All fields are required.", "error")
            return redirect(url_for("register"))

        if role == "admin":
            expected_code = os.getenv("ADMIN_REGISTER_CODE", "").strip()
            if not expected_code:
                flash("Admin registration is disabled until ADMIN_REGISTER_CODE is added in .env.", "error")
                return redirect(url_for("register"))
            if admin_code != expected_code:
                flash("Invalid admin access code.", "error")
                return redirect(url_for("register"))

        try:
            existing_user = users_collection.find_one({"email": email})
        except PyMongoError as error:
            app.logger.error("MongoDB registration connection error: %s", error)
            flash("Database connection issue. Please check MongoDB Atlas MONGO_URI and Network Access, then try again.", "error")
            return redirect(url_for("register"))
        if existing_user:
            flash("Email already registered. Please login.", "error")
            return redirect(url_for("login"))

        hashed_password = bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode(
            "utf-8"
        )

        users_collection.insert_one(
            {
                "name": name,
                "email": email,
                "password": hashed_password,
                "role": role,
                "quizzes_generated": 0,
                "codes_analyzed": 0,
                "topics_planned": 0,
                "chat_history": [],
                "created_at": datetime.utcnow(),
            }
        )

        flash("Registration successful. Please login.", "success")
        return redirect(url_for("login"))

    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        raw_email = request.form.get("email", "")
        try:
            email = normalize_and_validate_email(raw_email)
        except ValueError:
            flash("Invalid email or password.", "error")
            return redirect(url_for("login"))
        password = request.form.get("password", "")
        selected_role = request.form.get("role", "").strip().lower()

        try:
            user = users_collection.find_one({"email": email})
        except PyMongoError as error:
            app.logger.error("MongoDB login connection error: %s", error)
            flash("Database connection issue. Please check MongoDB Atlas MONGO_URI and Network Access, then try again.", "error")
            return redirect(url_for("login"))

        saved_password = user["password"].encode("utf-8") if user else b""

        if user and bcrypt.checkpw(password.encode("utf-8"), saved_password):
            session["user_id"] = str(user["_id"])
            session["name"] = user["name"]
            session["email"] = user["email"]
            session["role"] = user.get("role", "student")

            if selected_role in ["student", "admin"] and selected_role != session["role"]:
                flash(f"This account is registered as {session['role'].title()}, not {selected_role.title()}.", "error")
                session.clear()
                return redirect(url_for("login"))

            if session["role"] == "admin":
                return redirect(url_for("admin"))
            return redirect(url_for("dashboard"))

        flash("Invalid email or password.", "error")

    return render_template("login.html")


@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    if request.method == "POST":
        raw_email = request.form.get("email", "")
        try:
            email = normalize_and_validate_email(raw_email)
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("forgot_password"))

        user = users_collection.find_one({"email": email})
        if user:
            otp_code = f"{random.randint(100000, 999999)}"
            otp_hash = bcrypt.hashpw(otp_code.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")
            users_collection.update_one(
                {"_id": user["_id"]},
                {
                    "$set": {
                        "reset_otp_hash": otp_hash,
                        "reset_otp_expires": datetime.utcnow() + timedelta(minutes=10),
                    }
                },
            )
            sent, info = send_reset_email(email, otp_code)
            if not sent:
                flash(info, "error")
                return redirect(url_for("forgot_password"))

        session["reset_email"] = email
        flash("If that email exists, a 6-digit reset code has been sent.", "success")
        return redirect(url_for("verify_reset_code"))

    return render_template("forgot_password.html")


@app.route("/verify-reset-code", methods=["GET", "POST"])
def verify_reset_code():
    reset_email = session.get("reset_email", "")

    if request.method == "POST":
        raw_email = request.form.get("email", reset_email)
        try:
            email = normalize_and_validate_email(raw_email)
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("verify_reset_code"))
        code = request.form.get("code", "").strip()
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")

        if not email or not code or not new_password or not confirm_password:
            flash("Please fill all fields.", "error")
            return redirect(url_for("verify_reset_code"))

        if new_password != confirm_password:
            flash("Passwords do not match.", "error")
            return redirect(url_for("verify_reset_code"))

        user = users_collection.find_one({"email": email})
        valid = False
        if user and user.get("reset_otp_hash") and user.get("reset_otp_expires"):
            if datetime.utcnow() <= user["reset_otp_expires"]:
                valid = bcrypt.checkpw(code.encode("utf-8"), user["reset_otp_hash"].encode("utf-8"))

        if not valid:
            flash("Invalid or expired reset code.", "error")
            return redirect(url_for("verify_reset_code"))

        hashed_password = bcrypt.hashpw(new_password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")
        users_collection.update_one(
            {"_id": user["_id"]},
            {"$set": {"password": hashed_password}, "$unset": {"reset_otp_hash": "", "reset_otp_expires": ""}},
        )
        session.pop("reset_email", None)
        flash("Password reset successful. Please login.", "success")
        return redirect(url_for("login"))

    return render_template("verify_reset_code.html", reset_email=reset_email)


@app.route("/reset-password", methods=["POST"])
def reset_password():
    return redirect(url_for("verify_reset_code"))


@app.route("/logout")
def logout():
    session.clear()
    flash("Logged out successfully.", "success")
    return redirect(url_for("login"))


@app.route("/dashboard")
@role_required("student")
def dashboard():
    return render_template(
        "dashboard.html",
        name=session.get("name"),
    )


@app.route("/admin")
@role_required("admin")
def admin():
    total_users = users_collection.count_documents({})
    total_students = users_collection.count_documents({"role": "student"})
    return render_template(
        "admin.html",
        name=session.get("name"),
        total_users=total_users,
        total_students=total_students,
    )


@app.route("/feedback", methods=["GET", "POST"])
@login_required
def feedback():
    if request.method == "POST":
        subject = request.form.get("subject", "").strip()
        message = request.form.get("message", "").strip()

        if not subject or not message:
            flash("Please fill both subject and message.", "error")
            return redirect(url_for("feedback"))

        feedback_collection.insert_one(
            {
                "username": session.get("name", "Student"),
                "subject": subject,
                "message": message,
                "created_at": datetime.utcnow(),
            }
        )

        flash("Thanks! Your feedback has been sent to the admin.", "success")
        return redirect(url_for("feedback"))

    return render_template("feedback.html")






@app.route("/admin/feedback")
@role_required("admin")
def admin_feedback():
    feedback_messages = list(feedback_collection.find({}))
    feedback_messages.sort(
        key=lambda item: item.get("created_at", datetime.min),
        reverse=True,
    )

    return render_template(
        "admin_feedback.html",
        feedback_messages=feedback_messages,
    )





@app.route("/admin/users", methods=["GET", "POST"])
@role_required("admin")
def admin_users():
    if request.method == "POST":
        action = request.form.get("action")
        user_id = request.form.get("user_id")

        try:
            target_user_id = ObjectId(user_id)
        except Exception:
            flash("Invalid user selected.", "error")
            return redirect(url_for("admin_users"))

        if action == "update_role":
            new_role = request.form.get("role")

            if new_role not in ["student", "admin"]:
                flash("Invalid role selected.", "error")
                return redirect(url_for("admin_users"))

            users_collection.update_one(
                {"_id": target_user_id},
                {"$set": {"role": new_role}},
            )

            if str(target_user_id) == session.get("user_id"):
                session["role"] = new_role
                if new_role != "admin":
                    flash("Your role was changed to student.", "success")
                    return redirect(url_for("dashboard"))

            flash("User role updated.", "success")

        elif action == "delete_user":
            if str(target_user_id) == session.get("user_id"):
                flash("You cannot delete your own admin account.", "error")
                return redirect(url_for("admin_users"))

            users_collection.delete_one({"_id": target_user_id})
            flash("User deleted.", "success")

        return redirect(url_for("admin_users"))

    users = list(
        users_collection.find(
            {},
            {
                "name": 1,
                "role": 1,
                "topics_planned": 1,
                "quizzes_generated": 1,
                "codes_analyzed": 1,
            },
        )
    )
    users.sort(key=lambda user: user.get("name", "").lower())

    return render_template(
        "admin_users.html",
        users=users,
        current_user_id=session.get("user_id"),
    )


@app.route("/leaderboard")
@login_required
def leaderboard():
    users = users_collection.find(
        {},
        {
            "name": 1,
            "topics_planned": 1,
            "quizzes_generated": 1,
            "codes_analyzed": 1,
        },
    )

    students = []

    for user in users:
        topics_planned = user.get("topics_planned", 0)
        quizzes_generated = user.get("quizzes_generated", 0)
        codes_analyzed = user.get("codes_analyzed", 0)
        total_score = topics_planned + quizzes_generated + codes_analyzed

        students.append(
            {
                "name": user.get("name", "Student"),
                "topics_planned": topics_planned,
                "quizzes_generated": quizzes_generated,
                "codes_analyzed": codes_analyzed,
                "total_score": total_score,
            }
        )

    students.sort(key=lambda student: student["total_score"], reverse=True)

    return render_template("leaderboard.html", students=students)
@app.route("/chat", methods=["GET", "POST"])
@login_required
def chat():
    user_email = session.get("email")

    user = users_collection.find_one({"email": user_email})
    if not user:
        flash("User not found. Please login again.", "error")
        return redirect(url_for("logout"))

    chat_conversations = user.get("chat_conversations", [])

    # Convert old flat chat_history into one conversation one time
    old_history = user.get("chat_history", [])
    if old_history and not chat_conversations:
        first_user_msg = next(
            (msg.get("text", "New Chat") for msg in old_history if msg.get("role") == "user"),
            "New Chat",
        )

        migrated_chat = {
            "id": str(uuid4()),
            "title": first_user_msg[:45],
            "messages": old_history,
            "created_at": datetime.utcnow(),
            "updated_at": datetime.utcnow(),
        }

        chat_conversations = [migrated_chat]

        users_collection.update_one(
            {"email": user_email},
            {
                "$set": {"chat_conversations": chat_conversations},
                "$unset": {"chat_history": ""},
            },
        )

    # New Chat must only open an empty chat screen.
    # It must NOT delete old chats and must NOT auto-open old chat.
    if request.args.get("new") == "1":
        session["active_chat_id"] = "__new__"
        chat_topics = sorted(
            chat_conversations,
            key=lambda item: (item.get("pinned", False), item.get("updated_at", datetime.min)),
            reverse=True,
        )
        return render_template(
            "chat.html",
            question="",
            chat_history=[],
            chat_topics=chat_topics,
            active_chat_id="__new__",
        )

    # User clicked one old chat topic from sidebar
    selected_chat_id = request.args.get("chat_id")
    if selected_chat_id:
        session["active_chat_id"] = selected_chat_id

    active_chat_id = session.get("active_chat_id", "__new__")

    active_chat = None
    if active_chat_id != "__new__":
        for conversation in chat_conversations:
            if conversation.get("id") == active_chat_id:
                active_chat = conversation
                break

    if request.method == "POST":
        question = request.form.get("question", "").strip()
        form_chat_id = request.form.get("chat_id", "").strip()

        if form_chat_id:
            active_chat_id = form_chat_id
            session["active_chat_id"] = form_chat_id

        # Re-check active chat using form_chat_id
        active_chat = None
        if active_chat_id != "__new__":
            for conversation in chat_conversations:
                if conversation.get("id") == active_chat_id:
                    active_chat = conversation
                    break

        uploaded_file = request.files.get("chat_file")
        has_file = bool(uploaded_file and uploaded_file.filename)

        if not question and not has_file:
            flash("Please type a question or attach a file/image.", "error")
            return redirect(url_for("chat"))

        # Only the FIRST message after New Chat creates a sidebar topic.
        # Second/third/fourth messages use this same active_chat id.
        if active_chat is None:
            active_chat = {
                "id": str(uuid4()),
                "title": (question or uploaded_file.filename or "Uploaded material")[:45],
                "messages": [],
                "created_at": datetime.utcnow(),
                "updated_at": datetime.utcnow(),
            }
            chat_conversations.insert(0, active_chat)
            session["active_chat_id"] = active_chat["id"]

        uploaded_context = ""
        uploaded_name = ""
        if has_file:
            uploaded_name = uploaded_file.filename
            try:
                uploaded_context = extract_text_from_file(uploaded_file).strip()
                if not uploaded_context:
                    flash("I could not find readable text in that file/image. Try a clearer image or text-based file.", "error")
                    return redirect(url_for("chat", chat_id=active_chat["id"]))
            except Exception as error:
                flash(f"Could not read uploaded file: {error}", "error")
                return redirect(url_for("chat", chat_id=active_chat["id"]))

        try:
            answer = generate_chat_answer(active_chat, question, uploaded_context, uploaded_name)
        except Exception as error:
            answer = friendly_ai_error(error)

        user_message = {"role": "user", "text": question or f"Please explain the uploaded file: {uploaded_name}"}
        if uploaded_name:
            user_message["attachment_name"] = uploaded_name
            user_message["attachment_context"] = uploaded_context[:15000]
        active_chat["messages"].append(user_message)
        active_chat["messages"].append({"role": "assistant", "text": answer})
        active_chat["updated_at"] = datetime.utcnow()

        # Title remains first message. Do not create title for every message.
        if not active_chat.get("title"):
            active_chat["title"] = question[:45]

        users_collection.update_one(
            {"email": user_email},
            {
                "$set": {"chat_conversations": chat_conversations},
                "$unset": {"chat_history": ""},
            },
        )

        return redirect(url_for("chat", chat_id=active_chat["id"]))

    current_messages = active_chat.get("messages", []) if active_chat else []

    search_query = request.args.get("q", "").strip().lower()
    if search_query:
        chat_conversations = [
            chat for chat in chat_conversations
            if search_query in chat.get("title", "").lower()
            or any(search_query in msg.get("text", "").lower() for msg in chat.get("messages", []))
        ]

    chat_topics = sorted(
        chat_conversations,
        key=lambda item: (item.get("pinned", False), item.get("updated_at", datetime.min)),
        reverse=True,
    )

    return render_template(
        "chat.html",
        question="",
        chat_history=current_messages,
        chat_topics=chat_topics,
        active_chat_id=active_chat_id,
        search_query=request.args.get("q", ""),
    )


@app.route("/chat/delete/<chat_id>", methods=["POST"])
@login_required
def delete_chat(chat_id):
    user_email = session.get("email")
    user = users_collection.find_one({"email": user_email})
    chat_conversations = user.get("chat_conversations", []) if user else []
    chat_conversations = [chat for chat in chat_conversations if chat.get("id") != chat_id]
    users_collection.update_one({"email": user_email}, {"$set": {"chat_conversations": chat_conversations}})
    if session.get("active_chat_id") == chat_id:
        session["active_chat_id"] = "__new__"
    flash("Chat deleted successfully.", "success")
    return redirect(url_for("chat", new=1))


@app.route("/chat/rename/<chat_id>", methods=["POST"])
@login_required
def rename_chat(chat_id):
    new_title = request.form.get("title", "").strip()[:60]
    user_email = session.get("email")
    user = users_collection.find_one({"email": user_email})
    chat_conversations = user.get("chat_conversations", []) if user else []
    for chat_item in chat_conversations:
        if chat_item.get("id") == chat_id and new_title:
            chat_item["title"] = new_title
            chat_item["updated_at"] = datetime.utcnow()
            break
    users_collection.update_one({"email": user_email}, {"$set": {"chat_conversations": chat_conversations}})
    flash("Chat renamed successfully.", "success")
    return redirect(url_for("chat", chat_id=chat_id))


@app.route("/chat/pin/<chat_id>", methods=["POST"])
@login_required
def pin_chat(chat_id):
    user_email = session.get("email")
    user, chat_conversations, active_chat, _ = find_user_chat(user_email, chat_id)
    if not active_chat:
        flash("Chat not found.", "error")
        return redirect(url_for("chat", new=1))
    active_chat["pinned"] = not active_chat.get("pinned", False)
    active_chat["updated_at"] = datetime.utcnow()
    users_collection.update_one({"email": user_email}, {"$set": {"chat_conversations": chat_conversations}})
    flash("Chat pin updated.", "success")
    return redirect(url_for("chat", chat_id=chat_id))


@app.route("/chat/clear/<chat_id>", methods=["POST"])
@login_required
def clear_chat(chat_id):
    user_email = session.get("email")
    user, chat_conversations, active_chat, _ = find_user_chat(user_email, chat_id)
    if not active_chat:
        flash("Chat not found.", "error")
        return redirect(url_for("chat", new=1))
    active_chat["messages"] = []
    active_chat["updated_at"] = datetime.utcnow()
    users_collection.update_one({"email": user_email}, {"$set": {"chat_conversations": chat_conversations}})
    flash("Current chat cleared.", "success")
    return redirect(url_for("chat", chat_id=chat_id))


@app.route("/chat/regenerate/<chat_id>", methods=["POST"])
@login_required
def regenerate_chat(chat_id):
    user_email = session.get("email")
    user, chat_conversations, active_chat, _ = find_user_chat(user_email, chat_id)
    if not active_chat:
        flash("Chat not found.", "error")
        return redirect(url_for("chat", new=1))

    messages = active_chat.get("messages", [])
    last_user_text = ""
    last_user_context = ""
    last_user_file = ""
    last_assistant_index = None
    for index in range(len(messages) - 1, -1, -1):
        if messages[index].get("role") == "assistant" and last_assistant_index is None:
            last_assistant_index = index
        if messages[index].get("role") == "user":
            last_user_text = messages[index].get("text", "")
            last_user_context = messages[index].get("attachment_context", "")
            last_user_file = messages[index].get("attachment_name", "")
            break

    if not last_user_text:
        flash("No question found to regenerate.", "error")
        return redirect(url_for("chat", chat_id=chat_id))

    try:
        temp_chat = {**active_chat, "messages": messages[:index]}
        new_answer = generate_chat_answer(temp_chat, last_user_text, last_user_context, last_user_file)
    except Exception as error:
        new_answer = friendly_ai_error(error)

    if last_assistant_index is not None:
        messages[last_assistant_index]["text"] = new_answer
    else:
        messages.append({"role": "assistant", "text": new_answer})
    active_chat["updated_at"] = datetime.utcnow()
    users_collection.update_one({"email": user_email}, {"$set": {"chat_conversations": chat_conversations}})
    flash("Response regenerated.", "success")
    return redirect(url_for("chat", chat_id=chat_id))


@app.route("/chat/edit/<chat_id>/<int:message_index>", methods=["POST"])
@login_required
def edit_chat_message(chat_id, message_index):
    new_text = request.form.get("message_text", "").strip()
    user_email = session.get("email")
    user, chat_conversations, active_chat, _ = find_user_chat(user_email, chat_id)
    if not active_chat or not new_text:
        flash("Could not edit message.", "error")
        return redirect(url_for("chat", chat_id=chat_id))

    messages = active_chat.get("messages", [])
    if message_index < 0 or message_index >= len(messages) or messages[message_index].get("role") != "user":
        flash("Only your messages can be edited.", "error")
        return redirect(url_for("chat", chat_id=chat_id))

    previous_messages = messages[:message_index]
    existing_attachment_context = messages[message_index].get("attachment_context", "")
    existing_attachment_name = messages[message_index].get("attachment_name", "")
    active_chat["messages"] = previous_messages
    try:
        answer = generate_chat_answer(active_chat, new_text, existing_attachment_context, existing_attachment_name)
    except Exception as error:
        answer = friendly_ai_error(error)
    edited_user_message = {"role": "user", "text": new_text}
    if existing_attachment_name:
        edited_user_message["attachment_name"] = existing_attachment_name
        edited_user_message["attachment_context"] = existing_attachment_context
    active_chat["messages"] = previous_messages + [edited_user_message, {"role": "assistant", "text": answer}]
    active_chat["updated_at"] = datetime.utcnow()
    users_collection.update_one({"email": user_email}, {"$set": {"chat_conversations": chat_conversations}})
    flash("Message edited and answer updated.", "success")
    return redirect(url_for("chat", chat_id=chat_id))


@app.route("/chat/export/<chat_id>/<file_type>")
@login_required
def export_chat(chat_id, file_type):
    user_email = session.get("email")
    user, chat_conversations, active_chat, _ = find_user_chat(user_email, chat_id)
    if not active_chat:
        flash("Chat not found.", "error")
        return redirect(url_for("chat", new=1))
    lines = [active_chat.get("title", "Chat Export"), ""]
    for msg in active_chat.get("messages", []):
        speaker = "You" if msg.get("role") == "user" else "AI"
        attachment = f" [File: {msg.get('attachment_name')}]" if msg.get("attachment_name") else ""
        lines.append(f"{speaker}{attachment}:")
        lines.append(msg.get("text", ""))
        lines.append("")
    body = "\n".join(lines).strip()
    safe_title = (active_chat.get("title") or "chat_export")[:35].replace(" ", "_")
    if file_type == "pdf":
        return create_pdf_response("VFSTR Chat Export", body, safe_title)
    if file_type == "docx":
        return create_docx_response("VFSTR Chat Export", body, safe_title)
    flash("Invalid export format.", "error")
    return redirect(url_for("chat", chat_id=chat_id))


@app.route("/quiz", methods=["GET", "POST"])
@login_required
def quiz():
    import json
    import re

    study_material = ""
    question_count = ""
    quiz_items = []
    quiz_error = None

    user = users_collection.find_one({"email": session.get("email")})
    quizzes_generated = user.get("quizzes_generated", 0) if user else 0

    def clean_text(value):
        value = str(value or "").strip()
        value = re.sub(r"^\*+|\*+$", "", value).strip()
        return value

    def normalize_option_key(value):
        value = str(value or "").strip().upper()
        if value.startswith("A"):
            return "A"
        if value.startswith("B"):
            return "B"
        if value.startswith("C"):
            return "C"
        if value.startswith("D"):
            return "D"
        return value[:1]

    def normalize_quiz_items(raw_items):
        normalized = []
        for item in raw_items or []:
            if not isinstance(item, dict):
                continue
            options = item.get("options", {})
            if isinstance(options, list):
                options = {chr(65 + i): str(opt) for i, opt in enumerate(options[:4])}
            if not isinstance(options, dict):
                continue
            clean_options = {}
            for key in ["A", "B", "C", "D"]:
                option_value = options.get(key) or options.get(key.lower())
                if option_value:
                    clean_options[key] = clean_text(option_value)
            question = clean_text(item.get("question") or item.get("Question"))
            correct_answer = normalize_option_key(item.get("correct_answer") or item.get("answer") or item.get("Correct Answer"))
            explanation = clean_text(item.get("explanation") or item.get("Explanation") or "")
            if question and len(clean_options) == 4 and correct_answer in clean_options:
                normalized.append({
                    "question": question,
                    "options": clean_options,
                    "correct_answer": correct_answer,
                    "explanation": explanation or "Review the selected answer and compare it with the correct option.",
                })
        return normalized

    def parse_quiz_text(raw_text):
        text = str(raw_text or "").strip()
        if not text:
            return []
        # Remove markdown JSON fences if the model adds them.
        text = re.sub(r"^```(?:json)?", "", text.strip(), flags=re.IGNORECASE).strip()
        text = re.sub(r"```$", "", text.strip()).strip()
        try:
            data = json.loads(text)
            if isinstance(data, dict):
                data = data.get("questions") or data.get("quiz") or []
            parsed = normalize_quiz_items(data)
            if parsed:
                return parsed
        except Exception:
            pass

        blocks = re.split(r"(?=Question\s*\d+\s*[:.)-])", text, flags=re.IGNORECASE)
        parsed_items = []
        for block in blocks:
            block = block.strip()
            if not block or not re.search(r"Question\s*\d+", block, re.IGNORECASE):
                continue
            lines = [line.strip() for line in block.splitlines() if line.strip()]
            question_lines = []
            options = {}
            correct_answer = ""
            explanation_lines = []
            mode = "question"
            for line in lines:
                line = clean_text(line)
                if re.match(r"Question\s*\d+\s*[:.)-]", line, re.IGNORECASE):
                    line = re.sub(r"Question\s*\d+\s*[:.)-]\s*", "", line, flags=re.IGNORECASE).strip()
                    if line:
                        question_lines.append(line)
                    continue
                option_match = re.match(r"^([A-Da-d])\s*[.)-]\s*(.+)$", line)
                if option_match:
                    mode = "options"
                    options[option_match.group(1).upper()] = clean_text(option_match.group(2))
                    continue
                answer_match = re.match(r"^(?:Correct\s*)?Answer\s*[:\-]\s*([A-Da-d]).*", line, re.IGNORECASE)
                if answer_match:
                    mode = "answer"
                    correct_answer = answer_match.group(1).upper()
                    continue
                explanation_match = re.match(r"^Explanation\s*[:\-]\s*(.*)$", line, re.IGNORECASE)
                if explanation_match:
                    mode = "explanation"
                    if explanation_match.group(1).strip():
                        explanation_lines.append(clean_text(explanation_match.group(1)))
                    continue
                if mode == "question" and len(options) == 0:
                    question_lines.append(line)
                elif mode == "explanation":
                    explanation_lines.append(line)
            item = {
                "question": " ".join(question_lines),
                "options": options,
                "correct_answer": correct_answer,
                "explanation": " ".join(explanation_lines),
            }
            parsed_items.extend(normalize_quiz_items([item]))
        return parsed_items

    if request.method == "POST":
        study_material = request.form.get("study_material", "").strip()
        try:
            question_count = int(request.form.get("question_count", ""))
        except ValueError:
            flash("Please enter number of questions.", "error")
            return redirect(url_for("quiz"))

        if question_count < 3 or question_count > 10:
            flash("Please choose between 3 and 10 questions.", "error")
            return redirect(url_for("quiz"))

        if not study_material:
            flash("Please paste some study material first.", "error")
            return redirect(url_for("quiz"))

        try:
            prompt = f"""
Create exactly {question_count} multiple-choice questions from the study material below.
Return ONLY valid JSON. Do not add markdown, explanation outside JSON, or code fences.
Use this exact schema:
[
  {{
    "question": "Question text here",
    "options": {{"A": "Option A", "B": "Option B", "C": "Option C", "D": "Option D"}},
    "correct_answer": "A",
    "explanation": "Short explanation here"
  }}
]

Study material:
{study_material}
"""

            raw_quiz_text = ask_ai("You create clean, valid JSON MCQ quizzes for students.", prompt)
            quiz_items = parse_quiz_text(raw_quiz_text)
            if not quiz_items:
                quiz_error = "Quiz was generated, but it could not be converted into interactive options. Please try again."
            else:
                quiz_output_id = save_generated_output(
                    "quiz",
                    {
                        "quiz_items": quiz_items,
                        "study_material": study_material,
                        "question_count": question_count,
                    },
                    {"question_count": question_count},
                )
                session["latest_quiz_output_id"] = quiz_output_id
                users_collection.update_one({"email": session.get("email")}, {"$inc": {"quizzes_generated": 1}})
                quizzes_generated += 1
        except Exception as error:
            quiz_error = friendly_ai_error(error)
    else:
        saved_quiz = get_generated_content_from_session("latest_quiz_output_id", "quiz") or {}
        quiz_items = saved_quiz.get("quiz_items", [])
        study_material = saved_quiz.get("study_material", "")
        question_count = saved_quiz.get("question_count", "")

    return render_template(
        "quiz.html",
        study_material=study_material,
        question_count=question_count,
        quiz_items=quiz_items,
        quiz_error=quiz_error,
        quizzes_generated=quizzes_generated,
    )

@app.route("/question-generator", methods=["GET", "POST"])
@login_required
def question_generator():
    output_types = ["Questions & Answers", "Important Questions", "MCQs", "Viva Questions", "Interview Questions", "Short Notes"]
    difficulties = ["Easy", "Medium", "Hard"]
    output_type = ""
    difficulty = ""
    question_count = ""
    source_text = ""
    generated_text = None
    source_label = None

    if request.method == "POST":
        output_type = request.form.get("output_type", output_type)
        difficulty = request.form.get("difficulty", difficulty)
        uploaded_file = request.files.get("study_file")
        typed_text = request.form.get("source_text", "").strip()
        session.pop("latest_questions_output_id", None)

        try:
            question_count = int(request.form.get("question_count", ""))
        except ValueError:
            flash("Please enter item count.", "error")
            return redirect(url_for("question_generator"))

        if output_type not in output_types:
            flash("Please select output type.", "error")
            return redirect(url_for("question_generator"))
        if difficulty not in difficulties:
            flash("Please select difficulty.", "error")
            return redirect(url_for("question_generator"))
        if question_count < 3 or question_count > 30:
            flash("Please choose between 3 and 30 items.", "error")
            return redirect(url_for("question_generator"))

        if uploaded_file and uploaded_file.filename:
            try:
                source_text = extract_text_from_file(uploaded_file).strip()
            except Exception as error:
                flash(f"Could not extract text from the uploaded file: {error}", "error")
                return redirect(url_for("question_generator"))
        else:
            source_text = typed_text

        if not source_text:
            flash("Please upload study material or paste content first.", "error")
            return redirect(url_for("question_generator"))

        try:
            prompt = f"""
You are generating study content for a student.

Primary rule:
- First use the uploaded/user-provided material.
- If the material is incomplete for a useful response, you may add general AI knowledge.
- Clearly separate the response into these headings when needed:
  📄 Based on Uploaded Material
  🤖 Additional AI Knowledge
- Do not pretend unsupported details came from the uploaded file.

Generate: {output_type}
Difficulty: {difficulty}
Number of items: {question_count}

User-provided material:
{source_text}
"""
            generated_text = ask_groq("You create accurate academic questions, answers, notes, MCQs, and viva content.", prompt)
            output_id = save_generated_output(
                "questions",
                generated_text,
                {"output_type": output_type, "difficulty": difficulty, "question_count": question_count},
            )
            session["latest_questions_output_id"] = output_id
            source_label = "Uploaded/User Material + AI fallback if needed"
            question_generator_collection.insert_one({
                "user_email": session.get("email"),
                "username": session.get("name", "Student"),
                "output_type": output_type,
                "difficulty": difficulty,
                "question_count": question_count,
                "generated_text": generated_text,
                "created_at": datetime.utcnow(),
            })
        except Exception as error:
            generated_text = friendly_ai_error(error)

    return render_template(
        "question_generator.html",
        output_types=output_types,
        difficulties=difficulties,
        output_type=output_type,
        difficulty=difficulty,
        question_count=question_count,
        source_text=source_text,
        generated_text=generated_text,
        source_label=source_label,
    )


@app.route("/download-questions/<file_type>")
@login_required
def download_questions(file_type):
    generated_questions = get_generated_content_from_session("latest_questions_output_id", "questions")
    if not generated_questions:
        flash("Please generate questions before downloading.", "error")
        return redirect(url_for("question_generator"))
    if file_type == "docx":
        return create_docx_response("VFSTR Question Generator", generated_questions, "vfstr_questions")
    if file_type == "pdf":
        return create_pdf_response("VFSTR Question Generator", generated_questions, "vfstr_questions")
    flash("Invalid download format.", "error")
    return redirect(url_for("question_generator"))


@app.route("/debugger", methods=["GET", "POST"])
@login_required
def debugger():
    languages = ["Python", "Java", "C++", "JavaScript", "HTML/CSS"]
    actions = ["Debug & Fix Code", "Explain How Code Works"]
    language = ""
    action = ""
    code = ""
    result = None

    user = users_collection.find_one({"email": session.get("email")})
    codes_analyzed = user.get("codes_analyzed", 0) if user else 0

    if request.method == "POST":
        language = request.form.get("language", "")
        action = request.form.get("action", "")
        code = request.form.get("code", "").strip()

        if language not in languages:
            flash("Please select language.", "error")
            return redirect(url_for("debugger"))

        if action not in actions:
            flash("Please select action.", "error")
            return redirect(url_for("debugger"))

        if not code:
            flash("Please paste code first.", "error")
            return redirect(url_for("debugger"))

        try:
            prompt = f"""
Language: {language}
Action: {action}

Code:
{code}

If the action is "Debug & Fix Code":
1. Explain the problem in simple words.
2. Provide the corrected code.
3. Mention what changed.

If the action is "Explain How Code Works":
1. Explain the code step by step.
2. Mention important concepts.
3. Keep the explanation beginner-friendly.
"""

            result = ask_ai("You are a beginner-friendly code debugging assistant.", prompt)

            users_collection.update_one(
                {"email": session.get("email")},
                {"$inc": {"codes_analyzed": 1}},
            )
            codes_analyzed += 1
        except Exception as error:
            result = friendly_ai_error(error)

    return render_template(
        "debugger.html",
        languages=languages,
        actions=actions,
        language=language,
        action=action,
        code=code,
        result=result,
        codes_analyzed=codes_analyzed,
    )


@app.route("/planner", methods=["GET", "POST"])
@login_required
def planner():
    syllabus_topics = ""
    study_days = ""
    plan_text = None

    user = users_collection.find_one({"email": session.get("email")})
    topics_planned = user.get("topics_planned", 0) if user else 0

    if request.method == "POST":
        syllabus_topics = request.form.get("syllabus_topics", "").strip()

        try:
            study_days = int(request.form.get("study_days", ""))
        except ValueError:
            flash("Please enter study days.", "error")
            return redirect(url_for("planner"))

        if study_days < 1 or study_days > 30:
            flash("Please choose study days from 1 to 30.", "error")
            return redirect(url_for("planner"))

        if not syllabus_topics:
            flash("Please paste your syllabus or topics first.", "error")
            return redirect(url_for("planner"))

        try:
            prompt = f"""
Create a day-wise study plan for the syllabus/topics below.

Number of study days: {study_days}

Use this simple format:
Day 1:
- Topics to study:
- Practice task:
- Quick revision:

Keep the plan realistic and beginner-friendly.

Syllabus/topics:
{syllabus_topics}
"""

            plan_text = ask_ai("You create clear day-wise study plans for students.", prompt)
            output_id = save_generated_output(
                "plan",
                plan_text,
                {"study_days": study_days},
            )
            session["latest_plan_output_id"] = output_id

            users_collection.update_one(
                {"email": session.get("email")},
{"$inc": {"topics_planned": 1}},
            )
            topics_planned += 1
        except Exception as error:
            plan_text = friendly_ai_error(error)

    return render_template(
        "planner.html",
        syllabus_topics=syllabus_topics,
        study_days=study_days,
        plan_text=plan_text,
        topics_planned=topics_planned,
    )


@app.route("/download-plan/<file_type>")
@login_required
def download_plan(file_type):
    generated_plan = get_generated_content_from_session("latest_plan_output_id", "plan")
    if not generated_plan:
        flash("Please generate a study plan before downloading.", "error")
        return redirect(url_for("planner"))
    if file_type == "docx":
        return create_docx_response("VFSTR Study Plan", generated_plan, "vfstr_study_plan")
    if file_type == "pdf":
        return create_pdf_response("VFSTR Study Plan", generated_plan, "vfstr_study_plan")
    flash("Invalid download format.", "error")
    return redirect(url_for("planner"))


@app.route("/notes", methods=["GET", "POST"])
@login_required
def notes():
    notes_types = ["Short Notes", "Detailed Notes", "Exam Revision Notes"]
    topic_material = ""
    notes_type = ""
    notes_text = None

    if request.method == "POST":
        topic_material = request.form.get("topic_material", "").strip()
        notes_type = request.form.get("notes_type", "")
        uploaded_file = request.files.get("study_file")

        if notes_type not in notes_types:
            notes_type = ""

        if uploaded_file and uploaded_file.filename:
            try:
                topic_material = extract_text_from_file(uploaded_file).strip()
            except Exception:
                flash("Could not extract text from the uploaded file.", "error")
                return redirect(url_for("notes"))

            if not topic_material:
                flash("Could not extract text from the uploaded file.", "error")
                return redirect(url_for("notes"))

        if not topic_material:
            flash("Please paste text or upload PDF, DOCX, PPTX, TXT, or image file.", "error")
            return redirect(url_for("notes"))

        try:
            prompt = f"""
Create clean study notes from the topic or study material below.

Notes type: {notes_type}

If the notes type is Short Notes:
- Keep the notes brief.
- Use simple bullet points.

If the notes type is Detailed Notes:
- Explain the topic clearly.
- Include headings, definitions, and examples where useful.

If the notes type is Exam Revision Notes:
- Focus on key points, important terms, and likely exam reminders.
- Make it easy to revise quickly.

Topic or study material:
{topic_material}
"""

            notes_text = ask_ai("You create clean beginner-friendly study notes.", prompt)
            output_id = save_generated_output(
                "notes",
                notes_text,
                {"notes_type": notes_type},
            )
            session["latest_notes_output_id"] = output_id
        except Exception as error:
            notes_text = friendly_ai_error(error)

    return render_template(
        "notes.html",
        notes_types=notes_types,
        topic_material=topic_material,
        notes_type=notes_type,
        notes_text=notes_text,
    )


@app.route("/download-notes/<file_type>")
@login_required
def download_notes(file_type):
    generated_notes = get_generated_content_from_session("latest_notes_output_id", "notes")

    if not generated_notes:
        flash("Please generate notes before downloading.", "error")
        return redirect(url_for("notes"))

    if file_type == "docx":
        return create_docx_response("VFSTR Generated Notes", generated_notes, "vfstr_notes")
    if file_type == "pdf":
        return create_pdf_response("VFSTR Generated Notes", generated_notes, "vfstr_notes")
    flash("Invalid download format.", "error")
    return redirect(url_for("notes"))


@app.route("/assignment-generator", methods=["GET", "POST"])
@login_required
def assignment_generator():
    difficulties = ["Easy", "Medium", "Hard"]
    subject = ""
    topic = ""
    difficulty = ""
    question_count = ""
    assignment_text = None
    can_download_assignment = False

    if request.method == "POST":
        subject = request.form.get("subject", "").strip()
        topic = request.form.get("topic", "").strip()
        difficulty = request.form.get("difficulty", "")
        session.pop("latest_assignment_output_id", None)

        try:
            question_count = int(request.form.get("question_count", ""))
        except ValueError:
            flash("Please enter number of questions.", "error")
            return redirect(url_for("assignment_generator"))

        if difficulty not in difficulties:
            flash("Please select difficulty.", "error")
            return redirect(url_for("assignment_generator"))

        if question_count < 3 or question_count > 15:
            flash("Please choose between 3 and 15 questions.", "error")
            return redirect(url_for("assignment_generator"))

        if not subject or not topic:
            flash("Please enter both subject and topic.", "error")
            return redirect(url_for("assignment_generator"))

        try:
            prompt = f"""
Create a student assignment using the details below.

Subject: {subject}
Topic: {topic}
Difficulty: {difficulty}
Number of questions: {question_count}

Use this exact simple format:
Title:

Instructions:
- ...

Questions:
1. ...

Expected Learning Outcomes:
- ...

Keep the assignment clear, practical, and beginner-friendly.
"""

            assignment_text = ask_ai("You create useful academic assignments for students.", prompt)
            output_id = save_generated_output(
                "assignment",
                assignment_text,
                {"subject": subject, "topic": topic, "difficulty": difficulty, "question_count": question_count},
            )
            session["latest_assignment_output_id"] = output_id
            can_download_assignment = True
        except Exception as error:
            assignment_text = friendly_ai_error(error)

    return render_template(
        "assignment_generator.html",
        difficulties=difficulties,
        subject=subject,
        topic=topic,
        difficulty=difficulty,
        question_count=question_count,
        assignment_text=assignment_text,
        can_download_assignment=can_download_assignment,
    )


@app.route("/download-assignment/<file_type>")
@login_required
def download_assignment(file_type):
    generated_assignment = get_generated_content_from_session("latest_assignment_output_id", "assignment")

    if not generated_assignment:
        flash("Please generate an assignment before downloading.", "error")
        return redirect(url_for("assignment_generator"))

    if file_type == "docx":
        return create_docx_response("VFSTR Assignment", generated_assignment, "vfstr_assignment")
    if file_type == "pdf":
        return create_pdf_response("VFSTR Assignment", generated_assignment, "vfstr_assignment")
    flash("Invalid download format.", "error")
    return redirect(url_for("assignment_generator"))





@app.route("/analytics")
@login_required
def analytics():
    user = users_collection.find_one({"email": session.get("email")})

    topics_planned = user.get("topics_planned", 0) if user else 0
    quizzes_generated = user.get("quizzes_generated", 0) if user else 0
    codes_analyzed = user.get("codes_analyzed", 0) if user else 0
    report = None

    try:
        prompt = f"""
Create a short personalized study strategy report for this student.

Student name: {session.get("name")}
Topics planned: {topics_planned}
Quizzes generated: {quizzes_generated}
Codes analyzed: {codes_analyzed}

Keep the report simple, motivating, and beginner-friendly.
Include:
1. What the student is doing well.
2. What they should focus on next.
3. A short study strategy for the next few days.
"""

        report = ask_ai("You create short academic analytics reports for students.", prompt)
    except Exception as error:
        report = friendly_ai_error(error)

    return render_template(
        "analytics.html",
        name=session.get("name"),
        topics_planned=topics_planned,
        quizzes_generated=quizzes_generated,
        codes_analyzed=codes_analyzed,
        report=report,
    )


@app.route("/download-report")
@login_required
def download_report():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    user = users_collection.find_one({"email": session.get("email")})

    username = user.get("name", session.get("name", "Student")) if user else session.get("name", "Student")
    role = user.get("role", session.get("role", "student")) if user else session.get("role", "student")
    topics_planned = user.get("topics_planned", 0) if user else 0
    quizzes_generated = user.get("quizzes_generated", 0) if user else 0
    codes_analyzed = user.get("codes_analyzed", 0) if user else 0
    generated_date = datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")

    pdf_buffer = BytesIO()
    pdf = canvas.Canvas(pdf_buffer, pagesize=letter)

    pdf.setTitle("VFSTR Activity Report")
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(72, 740, "VFSTR Academic Command Center")
    pdf.drawString(72, 715, "Student Activity Report")

    pdf.setFont("Helvetica", 12)
    y_position = 675

    report_lines = [
        f"Username: {username}",
        f"Role: {role}",
        f"Topics Planned: {topics_planned}",
        f"Quizzes Generated: {quizzes_generated}",
        f"Codes Analyzed: {codes_analyzed}",
        f"Generated Date: {generated_date}",
    ]

    for line in report_lines:
        wrapped_lines = wrap(line, width=90) or [""]

        for wrapped_line in wrapped_lines:
            pdf.drawString(72, y_position, wrapped_line)
            y_position -= 20

        y_position -= 8

    pdf.save()
    pdf_buffer.seek(0)

    return Response(
        pdf_buffer.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": "attachment; filename=activity_report.pdf"},
    )


if __name__ == "__main__":
    app.run(debug=os.getenv("FLASK_DEBUG", "False").lower() == "true")
